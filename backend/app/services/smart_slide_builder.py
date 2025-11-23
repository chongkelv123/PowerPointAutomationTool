"""
Smart Slide Builder
Intelligently builds slides using template layouts and placeholders
"""
from pptx.util import Inches, Pt
from pptx.enum.shapes import PP_PLACEHOLDER
from typing import Dict, List, Optional
from app.services.template_analyzer import TemplateAnalyzer, adjust_font_size_for_overflow
import logging

logger = logging.getLogger(__name__)


def add_content_to_placeholder(placeholder, content_text: str, is_bullets: bool = True):
    """
    Add content to a placeholder with smart formatting

    Args:
        placeholder: The placeholder shape
        content_text: Text to add
        is_bullets: Whether to format as bullet points
    """
    if not placeholder or not placeholder.has_text_frame:
        return

    text_frame = placeholder.text_frame
    text_frame.clear()

    if is_bullets:
        # Add as bullet points
        lines = content_text.split('\n') if isinstance(content_text, str) else content_text
        for idx, line in enumerate(lines):
            if not line.strip():
                continue

            if idx == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()

            p.text = line.strip()
            p.level = 0

        # Adjust font size if needed
        adjust_font_size_for_overflow(text_frame, '\n'.join(str(l) for l in lines))
    else:
        # Add as plain text
        text_frame.text = content_text
        adjust_font_size_for_overflow(text_frame, content_text)


def place_image_in_placeholder(slide, placeholder, image_stream):
    """
    Place an image in a picture placeholder

    Args:
        slide: The slide object
        placeholder: The placeholder to fill
        image_stream: BytesIO stream of the image
    """
    try:
        # Get placeholder position and size
        left = placeholder.left
        top = placeholder.top
        width = placeholder.width
        height = placeholder.height

        # Remove the placeholder
        sp = placeholder.element
        sp.getparent().remove(sp)

        # Add the image in the same position
        slide.shapes.add_picture(image_stream, left, top, width=width, height=height)
    except Exception as e:
        logger.error(f"Error placing image in placeholder: {e}")


def add_slide_with_smart_layout(prs, analyzer: TemplateAnalyzer, slide_title: str,
                                 content_dict: dict, images: list = None) -> Optional[object]:
    """
    Add a slide using the most appropriate layout from the template

    Args:
        prs: Presentation object
        analyzer: TemplateAnalyzer instance
        slide_title: Title for the slide
        content_dict: Dictionary with 'text' and/or 'bullet_points'
        images: List of image data dictionaries

    Returns:
        The created slide object
    """
    # Determine layout requirements
    has_images = images and len(images) > 0
    has_content = bool(content_dict.get('text') or content_dict.get('bullet_points'))

    # Find best matching layout
    best_layout = analyzer.find_best_layout(
        has_title=True,
        has_content=has_content,
        has_images=has_images,
        num_columns=1
    )

    if not best_layout:
        # Fallback to first available layout
        best_layout = analyzer.layouts[0] if analyzer.layouts else None

    if not best_layout:
        logger.error("No suitable layout found")
        return None

    # Create slide with the selected layout
    layout = prs.slide_layouts[best_layout.index]
    slide = prs.slides.add_slide(layout)

    logger.info(f"Using layout '{best_layout.name}' (type: {best_layout.layout_type}) for slide '{slide_title}'")

    # Add title
    if slide.shapes.title:
        slide.shapes.title.text = slide_title
    else:
        # Find first text placeholder for title
        for shape in slide.shapes:
            if shape.is_placeholder and shape.has_text_frame:
                shape.text = slide_title
                break

    # Add content to appropriate placeholders
    if has_content:
        # Prepare content text
        bullet_points = content_dict.get('bullet_points', [])
        text_content = content_dict.get('text', '')

        if bullet_points:
            content_text = bullet_points if isinstance(bullet_points, list) else [bullet_points]
        else:
            content_text = text_content

        # Find content placeholder
        content_added = False
        for shape in slide.shapes:
            if shape.is_placeholder and shape.has_text_frame:
                placeholder_type = shape.placeholder_format.type

                # Skip title placeholders
                if placeholder_type in [PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE]:
                    continue

                # Use body or object placeholders for content
                if placeholder_type in [PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT]:
                    if isinstance(content_text, list):
                        add_content_to_placeholder(shape, content_text, is_bullets=True)
                    else:
                        add_content_to_placeholder(shape, content_text, is_bullets=False)
                    content_added = True
                    break

        # If no suitable placeholder found, add content manually
        if not content_added and (bullet_points or text_content):
            # Create text box for content
            left = Inches(0.5)
            top = Inches(1.5)
            width = Inches(9) if not has_images else Inches(4.5)
            height = Inches(5)

            text_box = slide.shapes.add_textbox(left, top, width, height)
            text_frame = text_box.text_frame
            text_frame.word_wrap = True

            if bullet_points:
                for idx, bullet in enumerate(bullet_points):
                    if idx == 0:
                        p = text_frame.paragraphs[0]
                    else:
                        p = text_frame.add_paragraph()
                    p.text = str(bullet)
                    p.level = 0
            else:
                text_frame.text = text_content

            adjust_font_size_for_overflow(text_frame, '\n'.join(str(b) for b in bullet_points) if bullet_points else text_content)

    # Add images
    if has_images:
        # Try to use picture placeholders first
        picture_placeholders = [
            shape for shape in slide.shapes
            if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        ]

        # Import necessary function
        from app.services.ppt_service import download_image, decode_base64_image

        for idx, img_data in enumerate(images):
            try:
                # Get image stream
                if img_data.get('url'):
                    img_stream = download_image(img_data['url'])
                elif img_data.get('base64'):
                    img_stream = decode_base64_image(img_data['base64'])
                else:
                    continue

                # Use placeholder if available
                if idx < len(picture_placeholders):
                    place_image_in_placeholder(slide, picture_placeholders[idx], img_stream)
                else:
                    # Add image manually
                    position = img_data.get('position', {})
                    left = Inches(position.get('left', 5.5))
                    top = Inches(position.get('top', 2))
                    width = Inches(position.get('width', 4))
                    height = Inches(position.get('height', 3.5))

                    slide.shapes.add_picture(img_stream, left, top, width=width, height=height)

            except Exception as e:
                logger.error(f"Error adding image: {e}")
                continue

    return slide


def add_title_slide(prs, analyzer: TemplateAnalyzer, title: str, subtitle: str = None, images: list = None):
    """
    Add a title slide using the template's title layout

    Args:
        prs: Presentation object
        analyzer: TemplateAnalyzer instance
        title: Main title text
        subtitle: Optional subtitle text
        images: Optional list of images for title slide

    Returns:
        The created slide object
    """
    # Get title layout
    title_layout_info = analyzer.get_title_layout()
    if not title_layout_info:
        title_layout_info = analyzer.layouts[0] if analyzer.layouts else None

    if not title_layout_info:
        logger.error("No title layout found")
        return None

    layout = prs.slide_layouts[title_layout_info.index]
    slide = prs.slides.add_slide(layout)

    # Set title
    if slide.shapes.title:
        slide.shapes.title.text = title

    # Set subtitle if provided
    if subtitle:
        for shape in slide.shapes:
            if shape.is_placeholder:
                placeholder_type = shape.placeholder_format.type
                if placeholder_type == PP_PLACEHOLDER.SUBTITLE:
                    shape.text = subtitle
                    break

    # Add images if provided
    if images:
        from app.services.ppt_service import add_image_to_slide
        for idx, img_data in enumerate(images):
            add_image_to_slide(
                slide,
                img_data,
                default_left=6 + (idx * 0.5),
                default_top=2
            )

    return slide
