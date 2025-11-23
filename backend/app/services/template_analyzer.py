"""
Template Analyzer Service
Analyzes PowerPoint templates to identify and map slide layouts intelligently
"""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Pt
from typing import Dict, List, Optional, Tuple
import logging

logger = logging.getLogger(__name__)


class LayoutInfo:
    """Information about a slide layout"""
    def __init__(self, layout_index: int, layout_name: str):
        self.index = layout_index
        self.name = layout_name
        self.has_title = False
        self.has_content = False
        self.has_subtitle = False
        self.content_placeholders = []
        self.image_placeholders = []
        self.text_placeholders = []
        self.layout_type = "unknown"

    def __repr__(self):
        return f"LayoutInfo(index={self.index}, name='{self.name}', type={self.layout_type})"


class TemplateAnalyzer:
    """Analyzes PowerPoint template structure"""

    def __init__(self, presentation: Presentation):
        self.prs = presentation
        self.layouts = []
        self.analyze_layouts()

    def analyze_layouts(self):
        """Analyze all slide layouts in the template"""
        for idx, layout in enumerate(self.prs.slide_layouts):
            layout_info = LayoutInfo(idx, layout.name)
            self._analyze_layout_placeholders(layout, layout_info)
            self._classify_layout_type(layout_info)
            self.layouts.append(layout_info)
            logger.info(f"Analyzed layout {idx}: {layout_info}")

    def _analyze_layout_placeholders(self, layout, layout_info: LayoutInfo):
        """Analyze placeholders in a layout"""
        for shape in layout.placeholders:
            placeholder_type = shape.placeholder_format.type

            # Check for title placeholder
            if placeholder_type == PP_PLACEHOLDER.TITLE or placeholder_type == PP_PLACEHOLDER.CENTER_TITLE:
                layout_info.has_title = True

            # Check for subtitle
            elif placeholder_type == PP_PLACEHOLDER.SUBTITLE:
                layout_info.has_subtitle = True

            # Check for body/content placeholders
            elif placeholder_type in [PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT]:
                layout_info.has_content = True
                layout_info.content_placeholders.append({
                    'index': shape.placeholder_format.idx,
                    'type': placeholder_type,
                    'left': shape.left,
                    'top': shape.top,
                    'width': shape.width,
                    'height': shape.height
                })

            # Check for picture placeholders
            elif placeholder_type == PP_PLACEHOLDER.PICTURE:
                layout_info.image_placeholders.append({
                    'index': shape.placeholder_format.idx,
                    'left': shape.left,
                    'top': shape.top,
                    'width': shape.width,
                    'height': shape.height
                })

            # Other text placeholders
            elif shape.has_text_frame:
                layout_info.text_placeholders.append({
                    'index': shape.placeholder_format.idx,
                    'type': placeholder_type,
                    'left': shape.left,
                    'top': shape.top,
                    'width': shape.width,
                    'height': shape.height
                })

    def _classify_layout_type(self, layout_info: LayoutInfo):
        """Classify the layout type based on its structure"""
        name_lower = layout_info.name.lower()

        # Title slide
        if layout_info.has_title and layout_info.has_subtitle and not layout_info.has_content:
            layout_info.layout_type = "title"
        # Title only
        elif layout_info.has_title and not layout_info.has_content and not layout_info.has_subtitle:
            layout_info.layout_type = "section_header"
        # Title with content
        elif layout_info.has_title and layout_info.has_content:
            if len(layout_info.content_placeholders) >= 2:
                layout_info.layout_type = "two_column"
            elif len(layout_info.image_placeholders) > 0:
                layout_info.layout_type = "content_with_image"
            else:
                layout_info.layout_type = "content"
        # Blank
        elif not layout_info.has_title and not layout_info.has_content:
            layout_info.layout_type = "blank"
        # Picture with caption
        elif len(layout_info.image_placeholders) > 0:
            layout_info.layout_type = "picture"
        else:
            # Fallback based on name
            if 'blank' in name_lower:
                layout_info.layout_type = "blank"
            elif 'title' in name_lower and 'only' in name_lower:
                layout_info.layout_type = "section_header"
            elif 'two' in name_lower or 'comparison' in name_lower or '2' in name_lower:
                layout_info.layout_type = "two_column"
            elif 'picture' in name_lower or 'image' in name_lower:
                layout_info.layout_type = "picture"
            elif 'content' in name_lower:
                layout_info.layout_type = "content"

    def find_best_layout(self, has_title: bool = True, has_content: bool = True,
                        has_images: bool = False, num_columns: int = 1) -> Optional[LayoutInfo]:
        """
        Find the best matching layout for the given requirements

        Args:
            has_title: Whether the slide needs a title
            has_content: Whether the slide has text content
            has_images: Whether the slide has images
            num_columns: Number of columns needed (1 or 2)

        Returns:
            LayoutInfo for the best matching layout, or None
        """
        candidates = []

        for layout in self.layouts:
            score = 0

            # Title slide for first slide
            if layout.layout_type == "title" and has_title and not has_content:
                score += 100

            # Content with images
            elif has_images and layout.layout_type == "content_with_image":
                score += 90

            # Two column layout
            elif num_columns >= 2 and layout.layout_type == "two_column":
                score += 85

            # Picture layout
            elif has_images and layout.layout_type == "picture":
                score += 80

            # Standard content
            elif has_content and layout.layout_type == "content":
                score += 75

            # Section header
            elif has_title and not has_content and layout.layout_type == "section_header":
                score += 70

            # Blank as fallback
            elif layout.layout_type == "blank":
                score += 10

            # Bonus points for matching features
            if has_title and layout.has_title:
                score += 5
            if has_content and layout.has_content:
                score += 5
            if has_images and len(layout.image_placeholders) > 0:
                score += 10

            if score > 0:
                candidates.append((score, layout))

        if candidates:
            # Return the highest scoring layout
            candidates.sort(key=lambda x: x[0], reverse=True)
            return candidates[0][1]

        # Fallback to first content layout
        for layout in self.layouts:
            if layout.layout_type == "content":
                return layout

        # Ultimate fallback - return any layout with content
        for layout in self.layouts:
            if layout.has_content or layout.has_title:
                return layout

        return None

    def get_title_layout(self) -> Optional[LayoutInfo]:
        """Get the title slide layout"""
        for layout in self.layouts:
            if layout.layout_type == "title":
                return layout
        # Fallback to first layout
        return self.layouts[0] if self.layouts else None

    def get_content_layout(self) -> Optional[LayoutInfo]:
        """Get a standard content layout"""
        for layout in self.layouts:
            if layout.layout_type == "content":
                return layout
        # Fallback
        return self.find_best_layout(has_title=True, has_content=True)

    def get_blank_layout(self) -> Optional[LayoutInfo]:
        """Get a blank layout"""
        for layout in self.layouts:
            if layout.layout_type == "blank":
                return layout
        # Return last layout as fallback
        return self.layouts[-1] if self.layouts else None


def analyze_template(template_path: str) -> TemplateAnalyzer:
    """
    Analyze a PowerPoint template

    Args:
        template_path: Path to the template file

    Returns:
        TemplateAnalyzer instance with layout information
    """
    prs = Presentation(template_path)
    return TemplateAnalyzer(prs)


def get_text_length(text: str) -> int:
    """
    Estimate the display length of text

    Args:
        text: The text to measure

    Returns:
        Estimated length score
    """
    if not text:
        return 0
    # Count characters, with extra weight for newlines
    return len(text) + text.count('\n') * 20


def adjust_font_size_for_overflow(text_frame, text: str, max_size: int = 24, min_size: int = 10):
    """
    Adjust font size to fit text in a text frame

    Args:
        text_frame: The text frame to adjust
        text: The text content
        max_size: Maximum font size in points
        min_size: Minimum font size in points
    """
    text_length = get_text_length(text)

    # Estimate appropriate font size based on text length
    if text_length < 200:
        font_size = max_size
    elif text_length < 400:
        font_size = max_size - 4
    elif text_length < 600:
        font_size = max_size - 6
    elif text_length < 800:
        font_size = max_size - 8
    else:
        font_size = min_size

    # Apply font size to all paragraphs
    for paragraph in text_frame.paragraphs:
        if paragraph.runs:
            for run in paragraph.runs:
                run.font.size = Pt(font_size)
        else:
            paragraph.font.size = Pt(font_size)


def should_split_content(text: str, threshold: int = 1000) -> bool:
    """
    Determine if content should be split across multiple slides

    Args:
        text: The text content
        threshold: Character threshold for splitting

    Returns:
        True if content should be split
    """
    return get_text_length(text) > threshold


def split_bullet_points(bullet_points: List[str], max_per_slide: int = 6) -> List[List[str]]:
    """
    Split bullet points across multiple slides

    Args:
        bullet_points: List of bullet points
        max_per_slide: Maximum bullets per slide

    Returns:
        List of bullet point groups
    """
    if len(bullet_points) <= max_per_slide:
        return [bullet_points]

    result = []
    for i in range(0, len(bullet_points), max_per_slide):
        result.append(bullet_points[i:i + max_per_slide])

    return result
