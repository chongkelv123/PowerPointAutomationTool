from pptx import Presentation
from pptx.util import Inches, Pt
import os
from datetime import datetime

async def create_presentation(title: str, slides: list[dict]) -> dict:
    """
    Create a PowerPoint presentation with the given title and slides

    Args:
        title: Presentation title
        slides: List of slide dictionaries with 'title' and 'content' keys

    Returns:
        Dictionary with presentation details
    """
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Add title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title

    # Add content slides
    for slide_data in slides:
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes

        title_shape = shapes.title
        body_shape = shapes.placeholders[1]

        title_shape.text = slide_data.get('title', 'Untitled Slide')

        if 'content' in slide_data:
            text_frame = body_shape.text_frame
            text_frame.text = slide_data['content']

    # Save presentation
    output_dir = "output"
    os.makedirs(output_dir, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{title.replace(' ', '_')}_{timestamp}.pptx"
    filepath = os.path.join(output_dir, filename)

    prs.save(filepath)

    return {
        "filename": filename,
        "filepath": filepath,
        "slide_count": len(prs.slides)
    }
